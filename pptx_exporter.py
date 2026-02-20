"""PPTX export module for FCC Coverage Map.

Generates a PowerPoint file with US state outlines and H3 hex coverage
shapes. Every element (each state, each hex) is an individual movable
freeform shape — fully editable in PowerPoint.
"""

import json
import math
import io
import os
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.dml.color import RGBColor
from lxml import etree

# ---- Slide dimensions: widescreen 13.333" x 7.5" ----
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# ---- Continental US bounds (lat/lng) ----
LAT_MIN = 24.0
LAT_MAX = 50.0
LNG_MIN = -125.0
LNG_MAX = -66.0

# ---- States to exclude from continental US ----
EXCLUDE_STATES = {
    "Alaska", "Hawaii", "Puerto Rico", "Guam",
    "U.S. Virgin Islands", "American Samoa",
    "Northern Mariana Islands",
}

# ---- Technology category mapping (mirrors data_processor_large.py) ----
TECH_CATEGORIES = {
    10: "copper",
    40: "cable",
    50: "fiber",
    60: "satellite",
    70: "wireless",
    71: "wireless",
    72: "wireless",
}


# ---- Web Mercator projection ----

def _lat_to_mercator_y(lat):
    """Convert latitude to Web Mercator Y value."""
    lat_rad = math.radians(lat)
    return math.log(math.tan(math.pi / 4 + lat_rad / 2))


# Pre-compute Mercator bounds
_MERC_Y_MIN = _lat_to_mercator_y(LAT_MIN)
_MERC_Y_MAX = _lat_to_mercator_y(LAT_MAX)
_MERC_Y_RANGE = _MERC_Y_MAX - _MERC_Y_MIN
_LNG_RANGE = LNG_MAX - LNG_MIN


def lnglat_to_slide(lng, lat):
    """Convert (lng, lat) to (x_emu, y_emu) on the slide.

    Uses Web Mercator projection so shapes align correctly.
    """
    x_frac = (lng - LNG_MIN) / _LNG_RANGE
    merc_y = _lat_to_mercator_y(lat)
    y_frac = 1.0 - (merc_y - _MERC_Y_MIN) / _MERC_Y_RANGE

    x_emu = int(x_frac * SLIDE_WIDTH)
    y_emu = int(y_frac * SLIDE_HEIGHT)
    return x_emu, y_emu


def _hex_color_to_rgb(hex_color):
    """Convert '#rrggbb' to pptx RGBColor."""
    hex_color = hex_color.lstrip("#")
    r = int(hex_color[0:2], 16)
    g = int(hex_color[2:4], 16)
    b = int(hex_color[4:6], 16)
    return RGBColor(r, g, b)


def _set_fill_transparency(shape, transparency_pct):
    """Set fill transparency on a shape via XML manipulation.

    python-pptx does not expose fill transparency directly.
    We add an <a:alpha> element to the solid fill's srgbClr.
    transparency_pct: 0 = opaque, 100 = fully transparent.
    """
    nsmap = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    sp_elem = shape._element
    srgb_elems = sp_elem.findall(".//a:srgbClr", nsmap)
    if srgb_elems:
        alpha_val = int((100 - transparency_pct) * 1000)
        alpha_elem = etree.SubElement(
            srgb_elems[0],
            "{http://schemas.openxmlformats.org/drawingml/2006/main}alpha",
        )
        alpha_elem.set("val", str(alpha_val))


def apply_filters(geojson, tech_filters, min_speed):
    """Filter GeoJSON features by technology category and min download speed.

    Mirrors the frontend's shouldShowFeature() logic.
    Returns a new FeatureCollection with only matching features.
    """
    filtered_features = []
    for feature in geojson.get("features", []):
        props = feature.get("properties", {})
        tech_code = props.get("technology", 0)
        tech_cat = TECH_CATEGORIES.get(tech_code, "other")

        if tech_cat not in tech_filters:
            continue
        if props.get("max_download", 0) < min_speed:
            continue

        filtered_features.append(feature)

    return {"type": "FeatureCollection", "features": filtered_features}


def _draw_freeform_polygon(slide, coords_lnglat, fill_rgb, line_rgb,
                           line_width=Emu(6350), transparency=0):
    """Draw a polygon as a Freeform shape on the slide.

    coords_lnglat: list of [lng, lat] pairs (GeoJSON polygon ring).
    fill_rgb: pptx RGBColor for fill.
    line_rgb: pptx RGBColor for border.
    line_width: border width in EMU.
    transparency: fill transparency percentage (0=opaque, 100=transparent).

    Returns the shape, or None if invalid.
    """
    # Convert all vertices to slide EMU coordinates
    points = []
    for lng, lat in coords_lnglat:
        x, y = lnglat_to_slide(lng, lat)
        points.append((x, y))

    if len(points) < 3:
        return None

    # Compute bounding box
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    left = min(xs)
    top = min(ys)
    width = max(xs) - left
    height = max(ys) - top

    if width <= 0 or height <= 0:
        return None

    # Build freeform: start pen at local origin (0, 0)
    freeform = slide.shapes.build_freeform(0, 0)

    # Move to first vertex (relative to bounding box origin)
    freeform.move_to(points[0][0] - left, points[0][1] - top)

    # Draw line segments to remaining vertices, close the shape
    remaining = [(px - left, py - top) for px, py in points[1:]]
    freeform.add_line_segments(remaining, close=True)

    # Convert: place shape at (left, top) on the slide
    shape = freeform.convert_to_shape(left, top)

    # Apply fill
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_rgb
    if transparency > 0:
        _set_fill_transparency(shape, transparency)

    # Apply border
    shape.line.color.rgb = line_rgb
    shape.line.width = line_width

    return shape


def _load_us_states():
    """Load US states GeoJSON from bundled data file."""
    data_path = os.path.join(
        os.path.dirname(__file__), "static", "data", "us_states.json"
    )
    with open(data_path, "r") as f:
        return json.load(f)


def _draw_us_states(slide):
    """Draw each continental US state as an individual freeform shape.

    Each state is a separate, movable shape in PowerPoint with a light
    grey fill and subtle border — matching the web map aesthetic.
    """
    geojson = _load_us_states()

    fill_rgb = RGBColor(0xE9, 0xEC, 0xEF)  # Light grey
    line_rgb = RGBColor(0xAD, 0xB5, 0xBD)  # Medium grey border

    for feature in geojson["features"]:
        name = feature["properties"].get("name", "")
        if name in EXCLUDE_STATES:
            continue

        geom = feature["geometry"]
        rings = []

        if geom["type"] == "Polygon":
            # Only use exterior ring (first ring)
            rings.append(geom["coordinates"][0])
        elif geom["type"] == "MultiPolygon":
            for polygon in geom["coordinates"]:
                rings.append(polygon[0])  # Exterior ring of each part

        for ring in rings:
            # Filter out points far outside bounds
            filtered_ring = [
                (lng, lat) for lng, lat in ring
                if LAT_MIN - 2 <= lat <= LAT_MAX + 2
                and LNG_MIN - 2 <= lng <= LNG_MAX + 2
            ]
            if len(filtered_ring) >= 3:
                _draw_freeform_polygon(
                    slide, filtered_ring,
                    fill_rgb=fill_rgb,
                    line_rgb=line_rgb,
                    line_width=Emu(9525),  # 0.75pt
                    transparency=0,
                )


def generate_pptx(provider_data_list, bg_image_path=None):
    """Generate a PPTX file with state outlines and hex coverage shapes.

    Every element is an individual, movable freeform shape:
    - Each US state outline is a separate shape
    - Each H3 hex is a separate shape

    Args:
        provider_data_list: list of dicts, each with:
            - "geojson": GeoJSON FeatureCollection (already filtered)
            - "color": hex color string like "#e74c3c"
            - "brand_name": provider name string
        bg_image_path: ignored (kept for API compatibility)

    Returns:
        io.BytesIO containing the PPTX file data.
    """
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    # Use blank layout
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Set slide background to light grey
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(0xF8, 0xF9, 0xFA)

    # Draw US states as individual shapes (bottom layer)
    _draw_us_states(slide)

    # Draw hexes for each provider (on top of states)
    for prov in provider_data_list:
        color_rgb = _hex_color_to_rgb(prov["color"])
        geojson = prov["geojson"]

        for feature in geojson.get("features", []):
            geometry = feature.get("geometry", {})
            if geometry.get("type") != "Polygon":
                continue

            coords = geometry.get("coordinates", [[]])[0]
            if len(coords) < 3:
                continue

            # Skip hexes outside continental US bounds
            lngs = [c[0] for c in coords]
            lats = [c[1] for c in coords]
            if max(lngs) < LNG_MIN or min(lngs) > LNG_MAX:
                continue
            if max(lats) < LAT_MIN or min(lats) > LAT_MAX:
                continue

            _draw_freeform_polygon(
                slide, coords,
                fill_rgb=color_rgb,
                line_rgb=color_rgb,
                line_width=Emu(6350),  # 0.5pt
                transparency=50,
            )

    # Save to BytesIO
    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
